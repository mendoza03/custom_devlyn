# -*- coding: utf-8 -*-

from odoo import _, fields, models
from odoo.exceptions import UserError
from datetime import date, datetime
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook
import base64
import zipfile
import io
import logging
import os

_logger = logging.getLogger(__name__)


class MasterBudgetImporter(models.TransientModel):
    _name = 'master.budget.importer'
    _description = 'Budget Importer'

    file = fields.Binary(string='File', required=True)
    filename = fields.Char(string='Filename')
    line_ids = fields.One2many(comodel_name='master.budget.importer.line', inverse_name='budget_importer_id', string='Lines')
    state = fields.Selection([
        ('get', 'Pending'),
        ('import', 'Import'),
        ('done', 'Done'),
    ], string='State', default='get')
    is_correct_to_import = fields.Boolean(string='Is correct to import', default=False)

    def action_read_files(self):
        if not self.file or not self.filename:
            raise UserError('No se ha cargado ningún archivo o falta el nombre del archivo.')
        file_extension = os.path.splitext(self.filename)[-1].lower()
        ALLOWED_EXTENSIONS = ['.zip', '.xlsm', '.xlsx']
        if file_extension not in ALLOWED_EXTENSIONS:
            raise UserError('Formato de archivo no compatible: %s. Compatible: %s' % (file_extension, ', '.join(ALLOWED_EXTENSIONS)))
        elif file_extension == '.zip':
            self._process_zip_file()
        elif file_extension in ['.xlsm', '.xlsx']:
            file_data = io.BytesIO(base64.b64decode(self.file))
            self._process_xlsm_file(file_data, self.filename)
        self.write({
            'state': 'import',
            'is_correct_to_import': all(line.state == 'done' for line in self.line_ids),
        })
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importador de presupuestos',
            'res_model': 'master.budget.importer',
            'res_id': self.id,
            'view_mode': 'form',
            'view_type': 'form',
            'target': 'new',
        }

    def action_import_data(self):
        budgets = []
        for record in self.line_ids:
            
            project = self.env['project.project'].create({
                'name': record.data['name'],
            })
            budget = self.env['master.budget'].create({
                'name': record.data['name'],
                'project_id': project.id,
            })
            budgets.append(budget.id)
            project_budget_line = False
            project_lines = [line for line in record.data['lines'] if line['type'] == 'project']
            for line in project_lines:
                project_budget_line = self.env['master.budget.line'].create({
                    'budget_id': budget.id,
                    'type': 'project',
                    'name': line['name'],
                    'code': line['code'],
                    'project_id': project.id,
                })
            task_lines = [line for line in record.data['lines'] if line['type'] == 'task']
            
            for line in task_lines:
                task = self.env['master.budget.line'].create({
                    'budget_id': budget.id,
                    'type': 'task',
                    'name': line['name'],
                    'code': line['code'],
                    'parent_id': project_budget_line.id if project_budget_line else False,
                })
                self.env['project.task'].create({
                    'project_id': project.id,
                    'master_budget_line_id': task.id,
                    'name': '%s - %s' % (task.name, task.code)
                })
            concept_lines = [line for line in record.data['lines'] if line['type'] == 'concept']
            for line in concept_lines:
                parent = self.env['master.budget.line'].search([('budget_id', '=', budget.id), ('type', '=', 'task'), ('name', '=', line['parent_name']), ('code', '=', line['parent_code'])], limit=1)
                self.env['master.budget.line'].create({
                    'budget_id': budget.id,
                    'type': 'concept',
                    'name': line['name'],
                    'code': line['code'],
                    'parent_id': parent.id,
                })
            budget_item_lines = [line for line in record.data['lines'] if line['type'] == 'budget_item']
            for line in budget_item_lines:
                parent = self.env['master.budget.line'].search([('budget_id', '=', budget.id), ('type', '=', 'concept'), ('name', '=', line['parent_name']), ('code', '=', line['parent_code'])], limit=1)
                self.env['master.budget.line'].create({
                    'budget_id': budget.id,
                    'type': 'budget_item',
                    'name': line['name'],
                    'code': line['code'],
                    'parent_id': parent.id,
                    'product_id': line['product_id'],
                    'price_unit': line['price_unit'],
                    'product_uom': line['product_uom'],
                    'quantity': line['quantity'],
                    'quantity_use_manual': line['quantity_use_manual'],
                    'amount_use_manual': line['amount_use_manual'],
                })
        if budgets:
            return {
                'name': 'Registros de presupuestos',
                'type': 'ir.actions.act_window',
                'res_model': 'master.budget',
                'view_mode': 'list,form',
                'domain': [('id', 'in', budgets)]
            }
        else:
            return

    def _process_xlsm_file(self, file_data, file_name):
        error_message = ''
        data = {
            'name': '',
            'lines': [],
        }
        try:
            workbook = load_workbook(file_data,  data_only=True, keep_vba=True)
            ws = workbook[workbook.sheetnames[0]]
            if 'Proyecto - Nombre' == ws['A1'].value:
                if isinstance(ws['A2'].value, str) and isinstance(ws['B2'].value, (str, int, float)):
                    name = '%s - %s' % (ws['A2'].value, ws['B2'].value)
                    if False:
                    # if self.env['project.project'].search([('name', '=', name)]):
                        _logger.info('2')
                        error_message += 'Proyecto: Ya existe un proyecto registrado con ese código y nombre.\n'
                        raise
                    else:
                        data['name'] = name
                        data['lines'].append({
                            'name': ws['A2'].value,
                            'code': ws['B2'].value,
                            'type': 'project',
                        })
                else:
                    error_message += 'Proyecto: El nombre o código no está asignado correctamente.\n'
                    raise
                tasks = set()
                concepts = {}
                row = 2
                while True:
                    name = ws.cell(row=row, column=3).value
                    code = ws.cell(row=row, column=4).value
                    if name is None and code is None:
                        break
                    elif name is None or code is None:
                        error_message += 'Tarea: El nombre o código no está asignado correctamente.\n'
                        raise
                    fullname = '%s - %s' % (name, code)
                    tasks.add(fullname)
                    row += 1
                for task in tasks:
                    data_task = task.split(' - ')
                    data['lines'].append({
                        'name': data_task[0],
                        'code': data_task[1],
                        'type': 'task',
                    })
                row = 2
                while True:
                    name = ws.cell(row=row, column=5).value
                    code = ws.cell(row=row, column=6).value
                    if name is None and code is None:
                        break
                    elif name is None or code is None:
                        error_message += 'Concepto: El nombre o código no está asignado correctamente.\n'
                        raise
                    fullname = '%s - %s' % (name, code)
                    task_name = ws.cell(row=row, column=3).value
                    task_code = ws.cell(row=row, column=4).value
                    if task_name is None or task_code is None:
                        error_message += 'Concepto: El registro (%s) no tiene una tarea asignada correctamente.\n' % fullname
                        break
                    if not fullname in concepts:
                        concepts[fullname] = {
                            'name': name,
                            'code': code,
                            'type': 'concept',
                            'parent_name': task_name,
                            'parent_code': task_code,
                        }
                    row += 1
                for concept in concepts:
                    data['lines'].append({
                        'name': concepts[concept]['name'],
                        'code': concepts[concept]['code'],
                        'type': 'concept',
                        'parent_name': concepts[concept]['parent_name'],
                        'parent_code': concepts[concept]['parent_code'],
                    })
                row = 1
                while True:
                    row += 1
                    name = ws.cell(row=row, column=7).value
                    if name is None:
                        break
                    concept_name = ws.cell(row=row, column=5).value
                    concept_code = ws.cell(row=row, column=6).value
                    if concept_name is None or concept_code is None:
                        error_message += 'Partida: El registro (%s) no tiene un concepto definido.\n' % name
                        continue
                    product = self.env['product.template'].search(['|', ('name', '=', name), ('default_code', '=', name)], limit=1)
                    if not product:
                        error_message += 'Partida: El producto (%s) no se encuentra registrado.\n' % (name)
                        continue
                    price_unit = ws.cell(row=row, column=8).value
                    if price_unit is None:
                        error_message += 'Partida: El precio unitario del producto (%s) no está definido.\n' % name
                        continue
                    uom = ws.cell(row=row, column=9).value
                    if uom is None:
                        error_message += 'Partida: La unidad de medida del producto (%s) no está definido.\n' % name
                        continue
                    uom = self.env['uom.uom'].search([('name', '=', uom)], limit=1)
                    if not uom:
                        error_message += 'Partida: La unidad de medida del producto (%s) no está registrado.\n' % name
                        continue
                    quantity = ws.cell(row=row, column=10).value
                    if quantity is None:
                        error_message += 'Partida: La cantidad del producto (%s) no está definido.\n' % name
                        continue
                    quantity_use_manual = ws.cell(row=row, column=11).value
                    if quantity_use_manual is None:
                        quantity_use_manual = 0
                    amount_use_manual = ws.cell(row=row, column=12).value
                    if amount_use_manual is None:
                        amount_use_manual = 0
                    data['lines'].append({
                        'name': product.name,
                        'code': product.default_code,
                        'type': 'budget_item',
                        'parent_name': concept_name,
                        'parent_code': concept_code,
                        'product_id': product.id,
                        'price_unit': price_unit,
                        'product_uom': uom.id,
                        'quantity': quantity,
                        'quantity_use_manual': quantity_use_manual,
                        'amount_use_manual': amount_use_manual,
                    })
            else:
                error_message += 'Verificar los encabezados.\n'
        except Exception as e:
            error_message += 'Error: existe un problema al procesar el archivo: %s' % str(e)
        self._create_line(file_name, data, 'done' if not error_message else 'error', 'Archivo correcto' if not error_message else error_message)

    def _process_zip_file(self):
        try:
            file_data = io.BytesIO(base64.b64decode(self.file))
            with zipfile.ZipFile(file_data, 'r') as zip_ref:
                for file_name in zip_ref.namelist():
                    if file_name.endswith('/'):
                        continue
                    if file_name.endswith('.xlsm'):
                        with zip_ref.open(file_name) as xlsm_file:
                            self._process_xlsm_file(xlsm_file, file_name)
                    else:
                        self._create_line(file_name, {}, 'error', 'Omitir archivo no compatible: %s' % file_name)
                        continue
        except zipfile.BadZipFile:
            raise UserError('El archivo cargado no es un archivo ZIP válido.')
        except Exception as e:
            raise UserError('Error al procesar el archivo ZIP: %s' % str(e))

    def _serialize_data(self, data):
        if isinstance(data, dict):
            return {k: self._serialize_data(v) for k, v in data.items()}
        elif isinstance(data, list):
            return [self._serialize_data(v) for v in data]
        elif isinstance(data, (date, datetime)):
            return data.isoformat()
        return data

    def _create_line(self, filename='', data=False, status='done', message=False):
        serialized_data = self._serialize_data(data) if data else {}
        self.env['master.budget.importer.line'].create({
            'budget_importer_id': self.id,
            'filename': filename,
            'data': serialized_data,
            'state': status,
            'message': message,
        })