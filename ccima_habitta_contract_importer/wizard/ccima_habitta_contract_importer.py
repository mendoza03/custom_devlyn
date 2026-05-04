# -*- coding: utf-8 -*-

from odoo import _, fields, models
from odoo.exceptions import UserError
from datetime import date, datetime
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook
import base64
import zipfile
import io
import logging
import os

_logger = logging.getLogger(__name__)


class CcimaHabittaContractImporter(models.TransientModel):
    _name = 'ccima.habitta.contract.importer'
    _description = 'Contract Importer'

    def _default_income_account(self):
        return self.env['ir.config_parameter'].sudo().get_param('real_estate_bits.income_account')

    file = fields.Binary(string='File', required=True)
    filename = fields.Char(string='Filename')
    line_ids = fields.One2many('ccima.habitta.contract.importer.line', 'contract_importer_id', string='Lines')
    state = fields.Selection([
        ('get', 'Pending'),
        ('import', 'Import'),
        ('done', 'Done'),
    ], string='State', default='get')
    templates = fields.Char(string='Templates')
    is_correct_to_import = fields.Boolean(string='Is correct to import', default=False)
    account_income = fields.Many2one(comodel_name='account.account', string='Income Account', default=_default_income_account)
    advance_payment_journal_id = fields.Many2one(comodel_name='account.journal', string='Advance Payment Journal')
    moratorium_interest = fields.Float(string='Moratorium Interest', default=0.0)

    def action_read_files(self):
        if not self.file or not self.filename:
            raise UserError('No se ha cargado ningún archivo o falta el nombre del archivo.')
        file_extension = os.path.splitext(self.filename)[-1].lower()
        ALLOWED_EXTENSIONS = ['.zip', '.xlsm']
        if file_extension not in ALLOWED_EXTENSIONS:
            raise UserError('Formato de archivo no compatible: %s. Compatible: %s' % (file_extension, ', '.join(ALLOWED_EXTENSIONS)))
        elif file_extension == '.zip':
            self._process_zip_file()
        elif file_extension == '.xlsm':
            file_data = io.BytesIO(base64.b64decode(self.file))
            self._process_xlsm_file(file_data, self.filename)
        self.write({
            'state': 'import',
            'is_correct_to_import': all(line.state == 'done' for line in self.line_ids),
        })
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importador de contratos',
            'res_model': 'ccima.habitta.contract.importer',
            'res_id': self.id,
            'view_mode': 'form',
            'view_type': 'form',
            'target': 'new',
        }

    def action_import_data(self):
        contracts = []
        for line in self.line_ids:
            product = self.env['product.template'].search([('id', '=', line.data['property_id'])], limit=1)
            contract = self.env['property.contract'].create({
                'property_id': product.id,
                'date': line.data['date'],
                'date_to': line.data['date_to'],
                'partner_id': line.data['partner_id'],
                'user_id': line.data['user_id'],
                'pricing': line.data['pricing'],
                'date_payment': line.data['date_payment'],
                'advance_payment_type': line.data['advance_payment_type'],
                'advance_payment': line.data['advance_payment'],
                'insurance_fee': 0,
                'contract_type': 'is_ownership',
                'template_id': line.data['template_id'],
                'account_income': self.account_income.id,
                'advance_payment_journal_id': self.advance_payment_journal_id.id,
                'moratorium_interest': self.moratorium_interest,
            })
            for loan_line in line.data['lines']:
                self.env['loan.line'].create({
                    'contract_id': contract.id,
                    'date': loan_line['date'],
                    'name': loan_line['name'],
                    'amount': loan_line['amount'],
                    'amount_paid_manual': loan_line['amount_paid_manual'],
                    'interest': loan_line['interest'],
                    'initial_balance': loan_line['initial_balance'],
                    'moratorium_interest_manual': loan_line['moratorium_interest_manual'],
                    'final_balance_manual': loan_line['final_balance_manual'],
                    'interest_of_interest_manual': loan_line['interest_of_interest_manual'],
                    'interest_of_accrued_manual': loan_line['interest_of_accrued_manual'],
                    'amount_capital': loan_line['amount_capital'],
                    'amount_to_capital': loan_line['amount_to_capital'],
                    'payment_state_manual': loan_line['payment_state_manual'],
                    'payment_date_manual': loan_line['payment_date_manual'],
                })
            contract.action_confirm()
            try:
                contract.recalculate_sequence()
            except Exception as e:
                _logger.warning(e)
            contracts.append(contract.id)
        if contracts:
            return {
                'name': 'Contrato de venta',
                'type': 'ir.actions.act_window',
                'res_model': 'property.contract',
                'view_mode': 'list,form',
                'domain': [('id', 'in', contracts)]
            }
        else:
            return

    def _process_xlsm_file(self, file_data, file_name):
        error_message = ''
        data = {
            'date': False,
            'date_to': False,
            'partner_id': False,
            'user_id': False,
            'property_id': False,
            'pricing': 0,
            'date_payment': False,
            'advance_payment_type': 'amount',
            'advance_payment': 0,
            'lines': [],
            'contract_type': 'is_ownership',
            'template_id': False,
        }
        try:
            workbook = load_workbook(file_data,  data_only=True, keep_vba=True)
            sheet_name = 'TABULADOR'
            if sheet_name not in workbook.sheetnames:
                error_message = 'Hoja de trabajo: La hoja %s no existe. Hojas disponibles (%s)' % (sheet_name, ', '.join(workbook.sheetnames))
            else:
                ws = workbook[sheet_name]
                partner = self.env['res.partner'].search_read([('name', '=', ws['C13'].value)], ['id'])
                if partner:
                    data['partner_id'] = partner[0]['id']
                else:
                    error_message += 'Cliente: El contacto %s no se encuentra en el sistema.\n' % ws['C13'].value

                installments = 0
                if isinstance(ws['Y18'].value, int):
                    installments = ws['Y18'].value
                    template = self.env['installment.template'].search([('duration_month', '=', installments)])
                    if not template:
                        error_message += 'Plantilla de pago: La plantilla de pago a plazos a meses (%s) no se encuentra en el sistema ().' % installments
                    data['template_id'] = template.id
                else:
                    error_message += 'Plantilla de pago: no se puede obtener la cantidad de meses en la celda Y18.\n'

                if ws['Y7'].is_date:
                    data['date'] = ws['Y7'].value.isoformat()
                    data['date_payment'] = ws['Y7'].value.isoformat()
                    if installments:
                        data['date_to'] = (ws['Y7'].value + relativedelta(months=installments)).isoformat()
                else:
                    error_message += 'Fecha de inicio: La celda Y7 no contiene una fecha.\n'

                if ws['C14'].value:
                    property = self.env['product.template'].search_read([('name', '=', ws['C14'].value), ('is_property', '=', True)], ['id'])
                    if property:
                        data['property_id'] = property[0]['id']
                    else:
                        error_message += 'Propiedad: La propiedad %s no se encuentra en el sistema.\n' % ws['C14'].value
                else:
                    error_message += 'Propiedad: No se encuentra un valor en la celda C14.\n'

                if ws['C15'].value:
                    data['pricing'] = ws['C15'].value
                else:
                    error_message += 'Precio: No se encuentra un valor en la celda C15.\n'

                if ws['C16'].value:
                    data['advance_payment'] = ws['C16'].value
                else:
                    error_message += 'Enganche: No se encuentra un valor en la celda C16.\n'

                if 'PERIODO' == ws['A23'].value:
                    row = 24
                    while True:
                        cell_period = ws.cell(row=row, column=1).value
                        if cell_period is None or not isinstance(cell_period, (float, int)):
                            break
                        interest_value = ws.cell(row=row, column=7).value
                        if isinstance(interest_value, datetime):
                            interest_value = 0.00
                        elif isinstance(interest_value, (float, int)):
                            pass
                        else:
                            interest_value = 0.00
                        amount = ws.cell(row=row, column=3).value
                        amount_capital = ws.cell(row=row, column=8).value
                        amount_to_capital = 0
                        if amount_capital > amount:
                            amount_to_capital = amount_capital - amount
                            amount_capital = amount
                        line_data = {
                            'date': ws.cell(row=row, column=2).value.isoformat(),
                            'name': 'MENSUALIDAD %s' % cell_period,
                            'amount': amount,
                            'initial_balance': ws.cell(row=row, column=4).value,
                            'amount_paid_manual': ws.cell(row=row, column=5).value if ws.cell(row=row, column=6).is_date else False,
                            'final_balance_manual': ws.cell(row=row, column=9).value,
                            'moratorium_interest_manual': ws.cell(row=row, column=10).value,
                            'interest_of_interest_manual': ws.cell(row=row, column=11).value,
                            'interest_of_accrued_manual': ws.cell(row=row, column=12).value,
                            'interest': interest_value,
                            'amount_capital': amount_capital,
                            'amount_to_capital': amount_to_capital,
                            'payment_state_manual': 'paid' if ws.cell(row=row, column=6).is_date else 'not_paid',
                            'payment_date_manual': ws.cell(row=row, column=6).value.isoformat() if ws.cell(row=row, column=6).is_date else False,
                        }
                        data['lines'].append(line_data)
                        row += 1
                else:
                    error_message += 'Pagos: no se detecta la columna "PERIODO".\n'
        except Exception as e:
            error_message = 'Error: existe un problema al procesar el archivo: %s' % str(e)
        self._create_line(file_name, data, 'done' if not error_message else 'error', 'Archivo correcto' if not error_message else error_message)

    def _process_zip_file(self):
        try:
            file_data = io.BytesIO(base64.b64decode(self.file))
            with zipfile.ZipFile(file_data, 'r') as zip_ref:
                for file_name in zip_ref.namelist():
                    if file_name.endswith('/'):
                        continue
                    if file_name.endswith('.xlsm'):
                        with zip_ref.open(file_name) as xlsm_file:
                            self._process_xlsm_file(xlsm_file, file_name)
                    else:
                        self._create_line(file_name, {}, 'error', 'Omitir archivo no compatible: %s' % file_name)
                        continue
        except zipfile.BadZipFile:
            raise UserError('El archivo cargado no es un archivo ZIP válido.')
        except Exception as e:
            raise UserError('Error al procesar el archivo ZIP: %s' % str(e))

    def _serialize_data(self, data):
        if isinstance(data, dict):
            return {k: self._serialize_data(v) for k, v in data.items()}
        elif isinstance(data, list):
            return [self._serialize_data(v) for v in data]
        elif isinstance(data, (date, datetime)):
            return data.isoformat()
        return data

    def _create_line(self, filename='', data=False, status='done', message=False):
        serialized_data = self._serialize_data(data) if data else {}
        self.env['ccima.habitta.contract.importer.line'].create({
            'contract_importer_id': self.id,
            'filename': filename,
            'data': serialized_data,
            'state': status,
            'message': message,
        })